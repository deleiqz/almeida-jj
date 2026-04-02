import sys
import os

try:
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("python-docx not installed. Installing now...")
    os.system("pip install python-docx")
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create Document
doc = Document()

# Set ABNT Margins
section = doc.sections[0]
section.top_margin = Cm(3)
section.left_margin = Cm(3)
section.bottom_margin = Cm(2)
section.right_margin = Cm(2)

# Set ABNT Normal Style Font (Arial, 12pt)
style = doc.styles['Normal']
font = style.font
font.name = 'Arial'
font.size = Pt(12)

def add_paragraph(text, bold=False, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    
    # 1.5 Line Spacing ABNT
    p.paragraph_format.line_spacing = 1.5
    
    # Paragraph indentation ABNT (1.25cm for first line)
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    
    run = p.add_run(text)
    run.bold = bold
    return p

def add_heading(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(text)
    run.bold = True
    run.font.name = 'Arial'
    run.font.size = Pt(12)

# TITLE
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.line_spacing = 1.5
run = title.add_run("CONTRATO DE PRESTAÇÃO DE SERVIÇOS DE DESENVOLVIMENTO E LICENCIAMENTO DE SOFTWARE")
run.bold = True
run.font.name = 'Arial'
run.font.size = Pt(12)

doc.add_paragraph()

# PARTIES
add_paragraph("CONTRATANTE: [Razão Social da Empresa ou Nome Completo], CNPJ/CPF nº [Preencher], com sede em [Preencher Endereço Completo], neste ato representada por seu representante legal [Nome do Responsável], [Nacionalidade], [Estado Civil], [Profissão], portador(a) do RG nº [Preencher] e CPF nº [Preencher], doravante denominada CONTRATANTE.", bold=False)

add_paragraph("CONTRATADA: Vanderlei Dias Mendes, brasileiro, [Estado Civil], Desenvolvedor de Software, portador do CPF nº 072.501.549-74 e RG nº [Preencher], com sede na Av. Osvaldo Valle Cordeiro 1260, AP 91 Torre 2, São Paulo/SP, doravante denominada CONTRATADA.", bold=False)

add_paragraph("As partes acima identificadas celebram o presente Contrato de Prestação de Serviços de Desenvolvimento, Licenciamento e Manutenção de Software (SaaS), que se regerá pelas seguintes cláusulas e condições:")

# CLAUSE 1
add_heading("CLÁUSULA 1 - DO OBJETO")
add_paragraph("O presente contrato tem como objeto:")
add_paragraph("I - O desenvolvimento, parametrização e implementação do Sistema de Gestão Exclusivo para a academia Almeida JJ, englobando controle financeiro, gestão de alunos com CRUD completo, ficha de inscrição digital nativa, painel dashboard inteligente e servidor de disparos de WhatsApp integrado;")
add_paragraph("II - O licenciamento de uso do software para a CONTRATANTE;")
add_paragraph("III - A prestação contínua de serviços de hospedagem em nuvem, manutenção técnica e correções de retaguarda técnica (bugs).")

# CLAUSE 2
add_heading("CLÁUSULA 2 - DOS VALORES E FORMA DE PAGAMENTO")
add_paragraph("Pela prestação dos serviços e licenciamento, a CONTRATANTE pagará à CONTRATADA os seguintes valores:")
add_paragraph("2.1. Taxa de Implementação (Setup): O valor fixo e único de R$ 3.250,00 (Três mil, duzentos e cinquenta reais), referente à programação, banco de dados, design da plataforma, design da ficha de matrícula exclusiva e infraestrutura de base. Forma de pagamento: 50% (cinquenta por cento) no aceite deste contrato e 50% (cinquenta por cento) na aprovação final do sistema, via transação PIX.")
add_paragraph("2.2. Taxa de Manutenção e Hospedagem (Mensalidade): O valor mensal de R$ 200,00 (Duzentos reais), referente ao licenciamento contínuo de uso do software, despesas técnicas com servidores de banco de dados, servidores de interface e suporte técnico preventivo. Vencimento: Todo dia [Preencher Dia] de cada mês, via transação PIX.")

# CLAUSE 3
add_heading("CLÁUSULA 3 - DO PRAZO E DA RENOVAÇÃO")
add_paragraph("3.1. O presente contrato é celebrado pelo prazo determinado de 12 (doze) meses, contados a partir da data da entrega oficial do sistema e seu pleno funcionamento em ambiente de produção.")
add_paragraph("3.2. Após os 12 (doze) meses iniciais, o contrato será prorrogado automaticamente por prazo indeterminado, podendo a mensalidade de R$ 200,00 sofrer reajuste anual pautado pelo índice IGPM ou IPCA acumulado no período.")

# CLAUSE 4
add_heading("CLÁUSULA 4 - DO INADIMPLEMENTO E DAS PENALIDADES")
add_paragraph("4.1. Multa por Atraso: Em caso de atraso superior a 5 (cinco) dias no pagamento da respectiva mensalidade:")
add_paragraph("a) Será cobrada multa moratória de 2% (dois por cento) sobre o valor nominal da parcela devida;", indent=False)
add_paragraph("b) Incidirão juros de mora de 1% (um por cento) ao mês, calculados de forma pro rata die.", indent=False)
add_paragraph("4.2. Suspensão do Sistema: Caso a inadimplência perdure por tempo superior a 15 (quinze) dias corridos, a CONTRATADA reserva-se o direito de suspender o acesso da CONTRATANTE ao sistema eletrônico temporariamente, até a regularização integral dos débitos em aberto.")
add_paragraph("4.3. Multa Rescisória (Quebra de Fidelidade): Tendo em vista o investimento analítico e técnico inicial estruturado a longo prazo, caso a CONTRATANTE solicite formalmente o cancelamento e desativação do sistema antes de cumpridos os 12 (doze) meses estipulados na Cláusula 3, a CONTRATANTE incorrerá no pagamento de uma multa rescisória equivalente a 15% (quinze por cento) do valor total das mensalidades restantes para completar o período anual.")
add_paragraph("4.4. Rescisão sem Multa: Ultrapassado o período de 12 (doze) meses de vigência, a CONTRATANTE figurará livre para rescindir o presente instrumento a qualquer momento, isenta de multas rescisórias, bastando avisar com 30 (trinta) dias de antecedência.")
add_paragraph("4.5. Mau Uso e Quebra de Contrato: Em caso de infrações graves à Legislação Brasileira ou uso indevido do módulo de disparos de WhatsApp para o envio de mensagens caracterizadas como SPAM malicioso, a CONTRATADA poderá encerrar o contrato motivadamente, de maneira imediata, sem aviso prévio e isenta de devolução de quantias já pagas.")

# CLAUSE 5
add_heading("CLÁUSULA 5 - DOS DIREITOS AUTORAIS E PROPRIEDADE INTELECTUAL")
add_paragraph("5.1. Resta estritamente estabelecido que o código-fonte, a arquitetura relacional de banco de dados e as tecnologias agregadas (Design UX/UI, scripts automatizadores, etc.) desenvolvidas são de propriedade intelectual e patrimonial exclusiva da CONTRATADA.")
add_paragraph("5.2. A entrega delineada na Cláusula 1 engloba a cessão do Direito de Uso (Licenciamento Comercial de Software como Serviço) à CONTRATANTE. É expressamente vedado à CONTRATANTE copiar, revender, transferir a terceiros ou reivindicar registro autoral do software para si.")
add_paragraph("5.3. Em contradição, toda a massa de dados cadastrais (nomes, contatos, dados financeiros de alunos, assinaturas) colhida e alimentada dentro da plataforma é de inteira e exclusiva propriedade material da CONTRATANTE. Ocorrendo distrato contratual, a CONTRATADA obriga-se a exportar e fornecer tais dados eletrônicos limpos em formato de planilha, condicionando-se apenas à quitação de faturas então pendentes.")

# CLAUSE 6
add_heading("CLÁUSULA 6 - SUPORTE TÉCNICO E BOAS PRÁTICAS")
add_paragraph("6.1. A prestação contínua de Manutenção prevê garantia de engenharia contra falhas na execução funcional da ferramenta, bem como o custeio operacional do fornecimento dos servidores telemáticos vitais.")
add_paragraph("6.2. Acordo de Nível de Serviço (SLA): O amparo técnico fluirá de forma remota em dias úteis, no horário comercial (09h00 às 18h00 - Horário de Brasília). Manifestações que impactem a continuidade primária das operações (parada total do sistema) ostentarão tempo de primeira resposta não superior a 24 horas úteis.")
add_paragraph("6.3. Novos Módulos: Qualquer requisição técnica que abranja a idealização de características sistêmicas não previstas originariamente escapa ao orçamento aqui fixado, suscitando viabilidade e valoração em Orçamento Aditivo individual.")
add_paragraph("6.4. Isenção sobre Regulamentações de Terceiros (Meta/WhatsApp): A CONTRATADA concebe e fornece a ponte transacional entre o sistema e o dispositivo telefônico. No entanto, não exerce controle imperativo sobre regulamentos impostos pela empresa Meta. Por conseguinte, afasta-se de responsabilidade e ônus de ressarcimentos na eventualidade de o terminal celular da CONTRATANTE sofrer penalidades, interrupções ou banimentos causados por excesso volúmico de envios, denúncias de terceiros ou guinadas nas diretrizes globais daquela plataforma.")
add_paragraph("6.5. Boas Práticas Anti-SPAM: A fim de dirimir os bloqueios mencionados na subcláusula anterior, a CONTRATANTE se compromete a aquecer periodicamente seu número institucional, modular volumes extensos através de lotes espaçados cronologicamente, comunicar-se exclusivamente com alunos opt-in, praticar o diálogo humano bilateral corriqueiro e acatar pontualmente solicitações voluntárias de descadastramento proferidas por destinatários.")

# CLAUSE 7
add_heading("CLÁUSULA 7 - PROTEÇÃO DE DADOS (LGPD)")
add_paragraph("7.1. Conforme os ditames da Lei Geral de Proteção de Dados (Lei nº 13.709/2018), a CONTRATANTE desponta categoricamente como Controladora dos dados pessoais de seus alunos, absorvendo para si o dever cível de colacionar o livre e manifesto consentimento obreiro e pautar a finalidade adequada da coleta antes de importá-los ou armazená-los na plataforma.")
add_paragraph("7.2. A CONTRATADA perfaz a figura de mera Operadora, abrigando matrizes instrumentais para viabilizar o processamento dos dados por intermédio da Controladora, de sorte que se insere alheia aos teores meritórios das informações ali processadas e isenta de corrompimentos legais inerentes à aquisição ilícita daqueles dados.")

# CLAUSE 8
add_heading("CLÁUSULA 8 - DISPOSIÇÕES GERAIS E FORO")
add_paragraph("8.1. Por estarem assim alinhados aos justos e perfeitos termos de vontade, as partes firmam o presente contrato instrumental para que produza seus incontestes efeitos legais.")
add_paragraph("8.2. Fica eleito o foro da Comarca de São Paulo, Estado de São Paulo, para dirimir pendências, divergências ou controvérsias decorrentes das estipulações deste contrato, afastando-se qualquer outro, por melhor assentado estivesse.")

doc.add_paragraph()
doc.add_paragraph()
add_paragraph("São Paulo, ______ de __________________________ de 2026.", indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()

# Signatures
def add_signature(name, role, details):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(30)
    
    run = p.add_run("_____________________________________________________")
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1
    
    r1 = p2.add_run(name + "\n")
    r1.bold = True
    p2.add_run(role + "\n")
    p2.add_run(details)

add_signature("Vanderlei Dias Mendes", "(Contratada / Desenvolvedor)", "CPF: 072.501.549-74")
add_signature("Almeida JJ / [Nome do Representante Legal]", "(Contratante)", "CPF/CNPJ: [Preencher]")

doc.add_paragraph()
add_paragraph("TESTEMUNHAS:", bold=True, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER)

add_signature("", "", "Nome:\nCPF:")
add_signature("", "", "Nome:\nCPF:")

doc.save(r'C:\Users\Goku\Documents\AG\Contrato_AlmeidaJJ_ABNT.docx')
print("DOCX Gerado com sucesso!")
